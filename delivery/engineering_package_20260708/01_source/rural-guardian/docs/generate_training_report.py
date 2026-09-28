from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from pathlib import Path


ROOT = Path(r"c:\Users\50453\Desktop\rural-guardian\rural-guardian")
DOCS_DIR = ROOT / "docs"
ASSETS_DIR = ROOT / "frontend" / "src" / "assets" / "docs"
OUTPUT_PATH = DOCS_DIR / "天津市大学软件学院项目实训总结报告-乡村守护者智慧养老平台.docx"


def set_run_font(run, font_name="宋体", size=12, bold=False):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(size)
    run.bold = bold


def set_style_font(style, font_name="宋体", size=12, bold=False):
    style.font.name = font_name
    style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    style.font.size = Pt(size)
    style.font.bold = bold


def add_toc(paragraph):
    fld_simple = OxmlElement("w:fldSimple")
    fld_simple.set(qn("w:instr"), r'TOC \o "1-3" \h \z \u')
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "右键更新目录后显示完整内容"
    run.append(t)
    fld_simple.append(run)
    paragraph._p.append(fld_simple)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, "Times New Roman", 10)


def add_heading(document, text, level):
    p = document.add_paragraph()
    p.style = document.styles[f"Heading {level}"]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_run_font(run, "黑体", 16 - (level - 1) * 2, True)
    if level == 1:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
    else:
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(3)


def add_paragraph(document, text, first_line=2):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(first_line * 0.74)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    set_run_font(run, "宋体", 12)


def add_bullets(document, items):
    for item in items:
        p = document.add_paragraph(style="List Bullet")
        p.paragraph_format.line_spacing = 1.5
        run = p.add_run(item)
        set_run_font(run, "宋体", 12)


def add_table(document, rows):
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(value)
            set_run_font(run, "宋体", 10.5, r_idx == 0)
    document.add_paragraph()


def add_image_with_caption(document, image_path, caption):
    if image_path.exists():
        try:
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            run.add_picture(str(image_path), width=Cm(15.5))
            cap = document.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap_run = cap.add_run(caption)
            set_run_font(cap_run, "宋体", 10.5)
            return
        except Exception:
            pass

    fallback = document.add_paragraph()
    fallback.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fallback_run = fallback.add_run(f"{caption}（图示资源已在项目目录中保留，可在最终排版时替换）")
    set_run_font(fallback_run, "宋体", 10.5)


def build_cover(document):
    section = document.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.5)

    p1 = document.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(36)
    run = p1.add_run("天津市大学软件学院")
    set_run_font(run, "黑体", 22, True)

    p2 = document.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p2.add_run("项目实训总结报告")
    set_run_font(run, "黑体", 24, True)

    p3 = document.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(36)
    run = p3.add_run("乡村守护者智慧养老平台")
    set_run_font(run, "黑体", 20, True)

    p4 = document.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p4.add_run("Rural Guardian Smart Elderly Care Platform")
    set_run_font(run, "Times New Roman", 14)

    document.add_paragraph()
    document.add_paragraph()

    info_rows = [
        ["学院", "天津市大学软件学院"],
        ["专业", "____________________"],
        ["班级", "____________________"],
        ["姓名", "____________________"],
        ["学号", "____________________"],
        ["指导教师", "____________________"],
        ["完成日期", "2026 年 6 月"],
    ]
    table = document.add_table(rows=len(info_rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(info_rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(value)
            set_run_font(run, "宋体", 12, j == 0)

    document.add_paragraph()
    tip = document.add_paragraph()
    tip.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tip_run = tip.add_run("说明：提交前请补充个人信息。")
    set_run_font(tip_run, "宋体", 10.5)
    document.add_page_break()


def build_toc(document):
    add_heading(document, "目录", 1)
    p = document.add_paragraph()
    add_toc(p)
    document.add_page_break()


def build_body(document):
    add_heading(document, "一、项目概述", 1)
    add_heading(document, "1.1 项目背景", 2)
    add_paragraph(
        document,
        "随着我国老龄化进程不断加快，农村地区空巢老人、独居老人数量持续增长，传统养老模式在健康监测、应急响应、日常照护和服务协同方面存在明显短板。乡村地区信息化基础相对薄弱、服务人员不足、响应链条长，使得老人跌倒、突发疾病、走失、长时间无人照看等问题具有较高风险。因此，构建一个集健康监测、预警处置、服务协同、家属关爱与政府监管于一体的智慧养老平台，具有明显的现实意义和社会价值。"
    )
    add_paragraph(
        document,
        "“乡村守护者智慧养老平台”即是在这一背景下设计与实现的综合性项目。平台围绕“老人、家属、村级专员、服务商、政府管理员”五类角色展开，建立从数据采集到预警分析、从工单流转到服务评价、从管理驾驶舱到实时大屏展示的完整闭环，目标是形成一个面向乡村养老场景的数字化治理与服务平台。"
    )
    add_heading(document, "1.2 项目目标", 2)
    add_bullets(document, [
        "建设多角色协同的智慧养老平台，实现政府监管、村级执行、服务商交付、家属关爱和老人自服务的统一接入。",
        "实现老人健康数据采集、异常预警、应急响应、服务工单、审批流转和客服沟通等关键业务闭环。",
        "通过数据大屏和三维地图等可视化手段提升平台运营感知能力与指挥调度效率。",
        "预留 IoT 设备接入能力，支持智能手环、MQTT 数据流、WebSocket 实时消息等扩展。"
    ])
    add_heading(document, "1.3 项目成果概述", 2)
    add_paragraph(
        document,
        "项目已完成前后端一体化实现。前端采用 Vue 3、Pinia、Vue Router、Element Plus、ECharts 与 Three.js 技术栈，后端采用 Node.js、Express、MySQL 与 WebSocket 作为主要基础设施。系统已实现登录认证、老人信息管理、健康档案管理、预警管理、工单管理、审批中心、设备监控、政府端运营驾驶舱、大屏展示以及客服实时通信等核心能力。"
    )

    add_heading(document, "二、需求分析", 1)
    add_heading(document, "2.1 业务痛点分析", 2)
    add_bullets(document, [
        "老人分布分散，传统人工巡访成本高、频率低，难以及时发现异常情况。",
        "村级专员、家属、服务商与政府部门之间信息不对称，处理链路缺少统一平台支撑。",
        "纸质台账和电话沟通方式难以追踪告警、工单和审批过程，责任边界不清晰。",
        "现有养老平台多以单端或双端为主，缺乏乡村场景下多角色协同治理能力。",
        "农村养老场景对适老化、低成本、可复制部署的要求更高。"
    ])
    add_heading(document, "2.2 用户角色分析", 2)
    add_table(document, [
        ["角色", "核心职责", "主要使用场景"],
        ["政府管理员", "全局监管、数据查看、统筹协调、服务配置", "查看驾驶舱、大屏、审批与设备运行状态"],
        ["村级专员", "本村老人管理、预警接单、线下核查、工单派发", "处理老人异常、确认上门服务"],
        ["服务商", "接收工单、执行服务、反馈结果", "医疗、照护、陪伴等服务交付"],
        ["家属", "查看老人健康状态、接收通知、确认结果", "关注老人安全与服务进度"],
        ["老人", "查看个人信息、发起 SOS、接受服务", "适老化的基础自服务与求助"]
    ])
    add_heading(document, "2.3 功能需求分析", 2)
    add_bullets(document, [
        "认证与权限：支持账号密码登录、JWT 鉴权、路由权限控制与多角色访问隔离。",
        "老人管理：维护老人基础信息、村庄归属、健康状态、家属绑定关系。",
        "健康管理：记录心率、血氧、体温、步数、睡眠等健康指标并形成档案。",
        "预警管理：支持告警创建、接单、处置、家属确认、升级与流转追踪。",
        "工单管理：支持创建、派单、接单、执行、评价以及流转日志记录。",
        "审批管理：支持家属绑定、设备绑定、服务商认证等审批流程。",
        "客服沟通：通过 WebSocket 支持老人呼叫请求与客服实时消息通信。",
        "运营展示：支持驾驶舱首页、大屏态势展示、图表分析和三维地图可视化。"
    ])
    add_heading(document, "2.4 非功能需求分析", 2)
    add_bullets(document, [
        "安全性：登录认证、敏感字段脱敏、接口限流、安全响应头与角色权限控制。",
        "可用性：页面结构清晰、适老化交互、关键流程可视化，便于业务人员快速上手。",
        "可扩展性：预留 IoT、MQTT、AI、短信、语音、地图等扩展接口。",
        "可维护性：前后端模块化划分清晰，代码结构具备较好的可阅读性和可扩展性。",
        "部署性：支持传统部署与 Docker 部署，适合比赛演示与后续落地扩展。"
    ])

    add_heading(document, "三、系统总体设计", 1)
    add_heading(document, "3.1 总体架构设计", 2)
    add_paragraph(
        document,
        "系统采用前后端分离架构。前端负责多角色页面渲染、状态管理、图表展示和交互流程，后端负责认证授权、业务处理、数据库读写、实时消息推送以及 IoT 设备扩展能力接入。平台以 HTTP/REST 作为主要业务通信方式，以 WebSocket 承担客服消息与实时推送，以 MySQL 作为核心业务数据库，并通过可选 IoT 模块承接智能手环与 MQTT 数据流。"
    )
    add_image_with_caption(document, ASSETS_DIR / "system-architecture.jpg", "图 3-1 系统总体架构示意图")
    add_heading(document, "3.2 技术架构与选型", 2)
    add_table(document, [
        ["层次", "技术选型", "选型原因"],
        ["前端", "Vue 3 + Vite + Pinia + Vue Router", "组件化开发效率高，适合多角色后台系统构建"],
        ["UI 与可视化", "Element Plus + ECharts + Three.js", "兼顾表单后台、大屏可视化与三维展示"],
        ["后端", "Node.js + Express", "开发效率高，适合中小规模业务系统与实时通信扩展"],
        ["数据库", "MySQL", "结构化数据管理成熟，适合当前多表业务关系存储"],
        ["实时通信", "WebSocket", "满足客服消息、状态通知和实时互动需求"],
        ["IoT 扩展", "MQTT + 设备引擎", "便于接入智能手环与实时设备数据"]
    ])
    add_heading(document, "3.3 功能模块设计", 2)
    add_table(document, [
        ["模块", "子模块", "功能说明"],
        ["用户中心", "登录、鉴权、权限控制", "支持多角色登录和访问边界控制"],
        ["老人档案", "基础信息、家属绑定、健康信息", "集中管理服务对象与关系链"],
        ["预警中心", "预警查询、接单、处置、确认、升级", "形成完整闭环响应流程"],
        ["工单中心", "创建、派发、接单、执行、评价", "支撑养老服务资源调度"],
        ["审批中心", "家属绑定、设备绑定、服务商认证", "支撑业务流程规范化管理"],
        ["设备监控", "设备状态、接入数据、监控大屏", "为 IoT 扩展预留统一入口"],
        ["运营驾驶舱", "核心指标、趋势图、大屏与地图", "支撑管理层全局决策与态势感知"]
    ])
    add_heading(document, "3.4 数据库设计", 2)
    add_paragraph(
        document,
        "数据库围绕用户、老人、健康记录、预警记录、工单记录、设备信息、村庄信息和服务商信息等核心实体展开设计。系统通过用户表和角色字段构建权限体系，通过老人表与村庄、设备、健康、家属等表形成业务关联，通过预警流转与工单流转实现事件追踪。"
    )
    add_image_with_caption(document, ASSETS_DIR / "er-diagram.jpg", "图 3-2 主要数据实体关系图")
    add_table(document, [
        ["核心数据表", "主要作用"],
        ["sys_user", "保存平台用户账号、角色与权限关联信息"],
        ["elderly_info", "保存老人基础信息与归属关系"],
        ["health_record", "保存老人健康监测记录"],
        ["alert_record", "保存异常预警记录与处理状态"],
        ["service_order", "保存服务工单与执行过程"],
        ["device_info", "保存设备编号、绑定状态与运行信息"],
        ["village_info", "保存村庄信息和地理归属"],
        ["provider_info", "保存服务商基础信息和资质信息"]
    ])

    add_heading(document, "四、关键业务流程设计", 1)
    add_heading(document, "4.1 登录与认证流程", 2)
    add_paragraph(
        document,
        "用户在登录页输入账号密码后，前端通过 Axios 将请求发送至 /api/auth/login 接口，后端校验用户名和密码，认证通过后返回包含用户身份信息的 JWT Token。前端将 Token 持久化到本地，在后续请求中通过请求拦截器自动附加到 Authorization 请求头。路由守卫根据 Token 状态和用户角色控制页面访问范围，从而实现多角色权限管理。"
    )
    add_heading(document, "4.2 预警响应流程", 2)
    add_paragraph(
        document,
        "当老人出现健康异常、设备异常或人工上报事件时，系统会生成预警记录，初始状态为待处理。村级专员或政府管理员可以进入预警中心查看待办数据，专员接单后状态转为处理中，完成现场核查或电话确认后可提交处置结果。若需要家属确认，则状态转为待确认；若家属认可，则预警关闭；若家属提出异议，则预警重新进入待处理池，并继续流转。对于风险较高或处理复杂的场景，系统还支持预警升级。"
    )
    add_image_with_caption(document, ASSETS_DIR / "alert-flow.jpg", "图 4-1 预警响应流程图")
    add_heading(document, "4.3 工单流转流程", 2)
    add_paragraph(
        document,
        "在养老服务场景中，平台允许基于老人需求或预警结果创建工单。工单创建后由村级专员或政府管理员进行派发，服务商接单后进入执行阶段，执行完成后由发起方或家属进行评价，从而实现“创建—派单—接单—执行—评价”的闭环服务链路。这一流程使平台不仅具备事件管理能力，也具备服务落地能力。"
    )
    add_image_with_caption(document, ASSETS_DIR / "order-flow.jpg", "图 4-2 服务工单流转流程图")

    add_heading(document, "五、系统实现", 1)
    add_heading(document, "5.1 前端实现", 2)
    add_paragraph(
        document,
        "前端采用 Vue 3 作为基础框架，通过 Pinia 管理全局用户状态、菜单权限与页面共享数据，通过 Vue Router 实现多角色页面路由组织。项目采用后台壳层布局，结合页面上下文、工作台操作区和状态栏，构建出适合政府端、村级端、服务商端等后台场景的交互框架。"
    )
    add_paragraph(
        document,
        "在界面层面，系统广泛使用 Element Plus 进行表单、表格、弹窗、分页等业务组件构建，同时结合设计系统样式文件统一视觉语言。对于驾驶舱和大屏模块，则使用 ECharts 展示统计图表，使用 Three.js 支撑三维地图或三维视觉表现，以增强运营展示效果。"
    )
    add_heading(document, "5.2 后端实现", 2)
    add_paragraph(
        document,
        "后端基于 Express 构建 RESTful API，围绕认证、老人管理、健康管理、预警管理、工单管理、设备管理、审批中心、大屏展示、客服系统等模块组织路由。系统通过中间件实现 CORS 处理、请求体解析、安全响应头、速率限制和 JWT 鉴权，保证接口访问安全与稳定。"
    )
    add_paragraph(
        document,
        "在数据层，系统当前运行版本采用 MySQL 连接池进行数据库访问，兼顾结构化业务数据管理和实际部署需求。同时，项目中保留了从 SQLite 迁移过来的兼容层逻辑，使部分原有查询模式得以平滑迁移。"
    )
    add_heading(document, "5.3 实时通信与 IoT 扩展实现", 2)
    add_paragraph(
        document,
        "系统内置 WebSocket 服务，用于客服呼叫请求、客服接听、通话结束和聊天消息的双向推送。在 IoT 扩展方面，平台预留了设备引擎、MQTT Broker、MQTT Bridge 和串口桥接能力，可用于智能手环等设备的实时数据接入。为了保证核心业务链路稳定，当前系统默认关闭 IoT 模块自动启动，避免高频日志与模拟数据影响登录和接口响应。"
    )
    add_heading(document, "5.4 大屏与可视化实现", 2)
    add_paragraph(
        document,
        "政府端数据大屏是本项目的展示亮点之一。系统对大屏页面进行了整体重构，新增了刷新机制、更新时间提示、告警滚动区和指挥入口，并将三维地图从单纯的炫技展示转化为更具业务操作性的风险工作区。通过区县检索、排序、风险筛选和下钻导航，大屏不仅能够“看”，还能辅助“管”和“调度”。"
    )

    add_heading(document, "六、系统测试与运行结果", 1)
    add_heading(document, "6.1 测试环境", 2)
    add_table(document, [
        ["项目", "测试环境"],
        ["操作系统", "Windows 10/11 开发环境"],
        ["前端运行", "Vite 开发服务器，默认端口 3000"],
        ["后端运行", "Node.js + Express，默认端口 8080"],
        ["数据库", "MySQL 8.x"],
        ["浏览器", "Chrome / Edge"]
    ])
    add_heading(document, "6.2 功能测试结果", 2)
    add_table(document, [
        ["测试项", "测试内容", "结果"],
        ["登录认证", "政府管理员、村级专员等多角色登录", "通过"],
        ["老人管理", "老人信息查询、详情展示、列表访问", "通过"],
        ["预警流程", "接单、处理、家属确认、升级", "通过"],
        ["工单流程", "创建、派单、接单、评价", "通过"],
        ["审批流程", "家属绑定、设备绑定、服务商认证", "通过"],
        ["客服系统", "WebSocket 呼叫推送与聊天", "通过"],
        ["数据大屏", "图表展示、地图交互、告警滚动", "通过"]
    ])
    add_heading(document, "6.3 性能与稳定性表现", 2)
    add_paragraph(
        document,
        "通过对项目运行链路的调试与优化，系统当前能够较稳定地支撑前后端联调和业务演示。尤其在登录与基础接口链路上，项目通过关闭默认 IoT 自动启动、控制高频日志输出等方式，显著提升了后端响应稳定性。前端构建产物能够正常生成，主要页面可顺利完成演示。"
    )
    add_bullets(document, [
        "前端生产构建能够正常完成，说明项目已具备发布基础。",
        "后端健康检查接口可正常返回，说明服务启动链路稳定。",
        "登录接口与主要业务接口已完成联调验证。",
        "大屏、地图与复杂页面在演示环境中可正常访问。"
    ])
    add_heading(document, "6.4 运行成果说明", 2)
    add_paragraph(
        document,
        "项目已形成较完整的比赛展示形态，能够从登录页进入政府端首页、预警中心、工单管理、审批中心、设备监控和数据大屏等核心页面，并可使用默认测试账号进行演示。平台结构完整、模块清晰、展示效果较好，具备较强的汇报与答辩表现力。"
    )

    add_heading(document, "七、项目创新点", 1)
    add_heading(document, "7.1 五端协同治理闭环", 2)
    add_paragraph(
        document,
        "与传统仅关注单一服务端的养老系统不同，本项目构建了政府、村级专员、服务商、家属和老人五端协同模式，将监管、执行、交付、确认与自服务整合在一个平台之中，形成了更完整的乡村养老治理闭环。"
    )
    add_heading(document, "7.2 业务流程闭环化设计", 2)
    add_paragraph(
        document,
        "无论是预警处理还是服务工单，系统都强调流转记录、状态迁移、责任主体和结果确认，避免传统系统只重“信息展示”而忽略“过程闭环”的问题。这种流程化设计使平台更贴近真实业务管理场景。"
    )
    add_heading(document, "7.3 可视化与实用性结合", 2)
    add_paragraph(
        document,
        "项目在数据大屏和三维地图上并未停留于单纯展示层，而是将可视化能力与风险筛选、区县检索、运营指标和告警态势结合，增强了页面的业务实用性和决策支撑价值。"
    )
    add_heading(document, "7.4 IoT 能力预留与低成本落地潜力", 2)
    add_paragraph(
        document,
        "项目预留了 MQTT、设备引擎、串口桥接和智能手环接入能力，为后续扩展为软硬件一体化乡村养老平台提供了基础。同时，相较于大规模商用养老平台，该方案在部署与扩展上具有更强的灵活性和低成本复制潜力。"
    )

    add_heading(document, "八、总结与展望", 1)
    add_heading(document, "8.1 项目总结", 2)
    add_paragraph(
        document,
        "通过本次项目实训，完成了一个围绕乡村养老场景展开的综合性信息系统设计与实现。项目不仅覆盖了前端界面开发、后端接口设计、数据库建模、实时通信与可视化展示，还体现了从业务理解、信息架构设计到系统落地实现的完整过程。整体来看，系统已具备较好的展示性、完整性和可扩展性。"
    )
    add_heading(document, "8.2 当前不足", 2)
    add_bullets(document, [
        "部分历史文档与当前实际运行架构之间仍存在版本差异，需要进一步统一更新。",
        "IoT 与设备接入链路已预留，但真实硬件数据接入仍可继续深化。",
        "部分页面仍以演示数据和前端推导模型为主，后续可逐步切换为真实业务数据。",
        "项目在公开部署前仍需进一步完成环境变量脱敏、数据脱敏与运维脚本完善。"
    ])
    add_heading(document, "8.3 未来展望", 2)
    add_paragraph(
        document,
        "未来可从三个方向继续推进：第一，进一步完善硬件设备接入，提升平台在真实乡村养老场景下的数据采集能力；第二，增强 AI 分析、风险预测和服务推荐能力，提升平台智能化水平；第三，针对移动端、适老化与基层服务流程继续打磨，使平台从“可展示”升级为“可落地、可复制、可运营”的乡村智慧养老解决方案。"
    )

    add_heading(document, "附录：默认测试账号", 1)
    add_table(document, [
        ["角色", "账号", "密码"],
        ["政府管理员", "gov_admin", "admin123"],
        ["村级专员", "village_staff", "staff123"],
        ["服务商", "provider_user", "provider123"],
        ["家属", "family_user", "family123"],
        ["老人", "elderly_user", "elderly123"]
    ])


def main():
    document = Document()

    set_style_font(document.styles["Normal"], "宋体", 12, False)
    for level in (1, 2, 3):
        set_style_font(document.styles[f"Heading {level}"], "黑体", 16 - (level - 1) * 2, True)

    build_cover(document)
    build_toc(document)
    build_body(document)

    footer = document.sections[0].footer.paragraphs[0]
    add_page_number(footer)

    document.save(str(OUTPUT_PATH))
    print(f"generated: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
